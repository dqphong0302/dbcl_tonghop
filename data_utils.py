#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tiện ích đọc dữ liệu & kiểm tra (validate) cho phần mềm Tổng hợp Khảo sát.

Mục đích (phục vụ tiêu chí "Tính khoa học" của SKKN):
- Dò cột TỰ ĐỘNG theo TIÊU ĐỀ thay vì hard-code index (cột AK = 36, ô F5...).
  -> Bền vững khi cấu trúc file thay đổi, có fallback an toàn.
- VALIDATE dữ liệu đầu vào: phát hiện file lỗi, thiếu cột, tổng % bất thường.
  -> Tạo "báo cáo kiểm tra dữ liệu" minh bạch, đáng tin cậy.

Hỗ trợ cả .xls (xlrd) và .xlsx (openpyxl) qua một lớp đọc thống nhất.
"""

import os
import re
from html.parser import HTMLParser
from dataclasses import dataclass, field
from typing import List, Optional

# ───────────────────────────────────────────────────────────────
# LỚP ĐỌC THỐNG NHẤT .xls / .xlsx
# ───────────────────────────────────────────────────────────────


class _Cell:
    """Đối tượng ô tương thích xlrd: cho phép truy cập .value."""
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class SheetReader:
    """Bọc 1 sheet để truy cập đồng nhất bất kể .xls hay .xlsx.

    Tương thích ngược với API của xlrd: hỗ trợ cả cell(r,c).value,
    cell_value(r,c), row_values(r), nrows, ncols, name.
    """

    def __init__(self, name: str, rows: List[List]):
        self.name = name
        self._rows = rows
        self.nrows = len(rows)
        self.ncols = max((len(r) for r in rows), default=0)

    def cell_value(self, r: int, c: int):
        if 0 <= r < self.nrows:
            row = self._rows[r]
            if 0 <= c < len(row):
                return row[c]
        return ""

    def cell(self, r: int, c: int) -> "_Cell":
        """Tương thích xlrd: sheet.cell(r, c).value."""
        return _Cell(self.cell_value(r, c))

    def row_values(self, r: int) -> List:
        if 0 <= r < self.nrows:
            row = list(self._rows[r])
            if len(row) < self.ncols:
                row += [""] * (self.ncols - len(row))
            return row
        return [""] * self.ncols


class ExcelReadError(Exception):
    """Lỗi đọc file Excel, thông điệp tiếng Việt dễ hiểu cho người dùng."""


def _detect_format(filepath: str) -> str:
    """Nhận dạng định dạng theo NỘI DUNG file (không tin phần mở rộng):
    "xls" (OLE2), "xlsx" (ZIP), "html" (trang web đổi đuôi .xls)."""
    with open(filepath, "rb") as f:
        head = f.read(512)
    if not head:
        raise ExcelReadError("File rỗng (0 byte)")
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if head.startswith(b"PK"):
        return "xlsx"
    text = head.lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if text.startswith((b"<html", b"<!doctype", b"<table", b"<meta", b"<?xml")) or b"<table" in head.lower():
        return "html"
    raise ExcelReadError("Không phải file Excel hợp lệ (có thể là file văn bản hoặc đã bị đổi đuôi)")


def _html_number(text):
    """Số trong bảng HTML: "12", "3.5" và dấu phẩy thập phân kiểu Việt "1,5" → float; còn lại giữ chuỗi."""
    if re.fullmatch(r"-?\d+(\.\d+)?", text):
        return float(text)
    if re.fullmatch(r"-?\d+,\d+", text):
        return float(text.replace(",", "."))
    return text


class _HTMLTables(HTMLParser):
    """Đọc các bảng <table> trong file HTML (nhiều hệ thống xuất "Excel" dưới dạng HTML)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self._row, self._cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
            span = str(dict(attrs).get("colspan") or "1")
            self._span = int(span) if span.isdigit() else 1
        elif tag == "br" and self._cell is not None:
            self._cell.append("\n")

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            text = re.sub(r"[ \t\r\xa0]+", " ", "".join(self._cell)).strip()
            value = _html_number(text)
            self._row.append(value)
            self._row.extend([""] * (self._span - 1))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.tables[-1].append(self._row)
            self._row = None


class WorkbookReader:
    """Đọc workbook .xls/.xlsx (hoặc HTML đổi đuôi .xls) thành danh sách SheetReader."""

    def __init__(self, filepath: str):
        self.filepath = filepath
        self.sheets: List[SheetReader] = []
        if not os.path.exists(filepath):
            raise ExcelReadError("File không tồn tại")
        try:
            fmt = _detect_format(filepath)
            {"xls": self._load_xls, "xlsx": self._load_xlsx, "html": self._load_html}[fmt](filepath)
        except ExcelReadError:
            raise
        except PermissionError:
            raise ExcelReadError("Không có quyền đọc file (file đang bị khóa hoặc mở bởi chương trình khác)")
        except Exception as e:
            raise ExcelReadError(f"File Excel bị hỏng hoặc không đọc được ({e})") from e

    def _load_xls(self, filepath):
        import xlrd
        # logfile=devnull: bản .exe không có console, xlrd ghi cảnh báo ra stdout (None) sẽ lỗi
        with open(os.devnull, "w") as devnull:
            wb = xlrd.open_workbook(filepath, logfile=devnull)
            for sh in wb.sheets():
                rows = [sh.row_values(r) for r in range(sh.nrows)]
                self.sheets.append(SheetReader(sh.name, rows))

    def _load_xlsx(self, filepath):
        import openpyxl
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        try:
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(["" if v is None else v for v in row])
                self.sheets.append(SheetReader(ws.title, rows))
        finally:
            wb.close()

    def _load_html(self, filepath):
        raw = open(filepath, "rb").read()
        m = re.search(rb'charset=["\']?([\w-]+)', raw[:2048], re.I)
        text = None
        for enc in ([m.group(1).decode()] if m else []) + ["utf-8", "cp1258", "latin-1"]:
            try:
                text = raw.decode(enc)
                break
            except (LookupError, UnicodeDecodeError):
                continue
        parser = _HTMLTables()
        parser.feed(text)
        tables = [t for t in parser.tables if t]
        if not tables:
            raise ExcelReadError("File HTML không chứa bảng dữ liệu nào")
        for i, rows in enumerate(tables, 1):
            self.sheets.append(SheetReader(f"Sheet{i}", rows))

    @property
    def nsheets(self) -> int:
        return len(self.sheets)

    def sheet_names(self) -> List[str]:
        return [s.name for s in self.sheets]

    def sheet_by_index(self, idx: int) -> SheetReader:
        return self.sheets[idx]


def open_workbook(filepath: str) -> WorkbookReader:
    """Mở file Excel (.xls/.xlsx/HTML đổi đuôi) trả về WorkbookReader thống nhất.
    Lỗi đọc file → ExcelReadError với thông điệp tiếng Việt."""
    return WorkbookReader(filepath)


# ───────────────────────────────────────────────────────────────
# DÒ CỘT TỰ ĐỘNG THEO TIÊU ĐỀ
# ───────────────────────────────────────────────────────────────


def _norm(text) -> str:
    """Chuẩn hóa chuỗi để so khớp: bỏ khoảng trắng thừa, về chữ thường."""
    return re.sub(r"\s+", " ", str(text or "").strip().lower())


def find_column_by_keywords(
    sheet: SheetReader,
    keywords: List[str],
    max_scan_rows: int = 12,
) -> Optional[int]:
    """Tìm chỉ số CỘT đầu tiên mà tiêu đề chứa 1 trong các keyword.

    Quét tối đa `max_scan_rows` dòng đầu (vùng header). Trả về index cột hoặc None.
    """
    kws = [_norm(k) for k in keywords if _norm(k)]
    if not kws:
        return None
    for r in range(min(max_scan_rows, sheet.nrows)):
        for c in range(sheet.ncols):
            val = _norm(sheet.cell_value(r, c))
            if not val:
                continue
            if any(kw in val for kw in kws):
                return c
    return None


def find_header_row(
    sheet: SheetReader,
    keywords: List[str],
    max_scan_rows: int = 12,
) -> Optional[int]:
    """Tìm DÒNG header chứa 1 trong các keyword (vd dòng có 'STT', 'Họ đệm')."""
    kws = [_norm(k) for k in keywords if _norm(k)]
    for r in range(min(max_scan_rows, sheet.nrows)):
        row_text = " ".join(_norm(sheet.cell_value(r, c)) for c in range(sheet.ncols))
        if any(kw in row_text for kw in kws):
            return r
    return None


def detect_feedback_column(
    sheet: SheetReader,
    search_keywords: List[str],
) -> tuple:
    """Dò cột chứa ý kiến mở trong FILE FEEDBACK (raw, mỗi dòng 1 SV).

    Trả về (col_index, start_row).
    Chiến lược:
      1) Tìm theo keyword tiêu đề (vd 'ý kiến đóng góp').
      2) Fallback: cột cuối cùng, bắt đầu sau dòng header dò được (hoặc dòng 5).
    """
    col = find_column_by_keywords(sheet, search_keywords)
    if col is not None:
        # start_row = ngay sau dòng chứa tiêu đề
        header_row = find_header_row(sheet, ["stt", "họ đệm", "mã số sinh viên"])
        start = (header_row + 1) if header_row is not None else 5
        return col, start

    # Fallback an toàn
    header_row = find_header_row(sheet, ["stt", "họ đệm", "mã số sinh viên"])
    start = (header_row + 1) if header_row is not None else 5
    return sheet.ncols - 1, start


# ───────────────────────────────────────────────────────────────
# VALIDATE DỮ LIỆU
# ───────────────────────────────────────────────────────────────


@dataclass
class ValidationIssue:
    level: str            # "error" | "warning" | "info"
    location: str         # vd "Sheet 1", "C5"
    message: str


@dataclass
class ValidationReport:
    filepath: str
    issues: List[ValidationIssue] = field(default_factory=list)

    def add(self, level: str, location: str, message: str):
        self.issues.append(ValidationIssue(level, location, message))

    @property
    def errors(self):
        return [i for i in self.issues if i.level == "error"]

    @property
    def warnings(self):
        return [i for i in self.issues if i.level == "warning"]

    @property
    def ok(self) -> bool:
        return len(self.errors) == 0

    def summary(self) -> str:
        name = os.path.basename(self.filepath)
        return (f"{name}: {len(self.errors)} lỗi, {len(self.warnings)} cảnh báo, "
                f"{len(self.issues)} mục kiểm tra")


def validate_survey_sheet(sheet: SheetReader, sheet_label: str) -> ValidationReport:
    """Kiểm tra 1 sheet khảo sát (dạng đã tổng hợp, 8 cột).

    Kiểm tra: có vùng 'Thông tin khảo sát', tổng tỷ lệ ~ 100%, số cột tối thiểu.
    """
    report = ValidationReport(filepath=sheet_label)

    if sheet.nrows < 6:
        report.add("error", sheet_label, "Sheet quá ít dòng, không giống file khảo sát.")
        return report

    if sheet.ncols < 8:
        report.add("warning", sheet_label,
                   f"Chỉ có {sheet.ncols} cột (kỳ vọng >= 8). Có thể sai định dạng.")

    # Tìm vùng câu hỏi khảo sát
    has_survey = False
    for r in range(sheet.nrows):
        if _norm(sheet.cell_value(r, 0)) == "thông tin khảo sát":
            has_survey = True
            break
    if not has_survey:
        report.add("warning", sheet_label,
                   "Không tìm thấy mục 'Thông tin khảo sát'. Báo cáo có thể thiếu dữ liệu.")

    # Kiểm tra dòng 'Tỷ lệ' tổng ~ 1.0
    checked = 0
    for r in range(sheet.nrows):
        if _norm(sheet.cell_value(r, 1)) == "tỷ lệ":
            try:
                total = float(sheet.cell_value(r, 7))
            except (ValueError, TypeError):
                continue
            checked += 1
            if total > 0 and abs(total - 1.0) > 0.05:
                report.add("warning", f"{sheet_label} dòng {r+1}",
                           f"Tổng tỷ lệ = {total:.4f} (lệch khỏi 100%).")
    if checked:
        report.add("info", sheet_label, f"Đã kiểm tra {checked} dòng tỷ lệ.")

    return report
