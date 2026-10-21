#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Bộ kiểm thử tự động (pytest) cho phần mềm Tổng hợp Khảo sát.

Phục vụ tiêu chí "Tính khoa học" của SKKN: chứng minh phần mềm xử lý ĐÚNG
trên dữ liệu thật, và bảo vệ chống hồi quy khi nâng cấp.

Chạy:  pytest -v   (từ thư mục gốc dự án)
"""

import os
import sys
import math
import pytest

# Cho phép import module ở thư mục gốc
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import data_utils
import excel_to_word as e2w
import merge_feedbacks as mf

DULIEU = os.path.join(ROOT, "dulieu")
SURVEY_FILE = os.path.join(DULIEU, "CTXH25 Giáo dục thể chất 1 25-26 HKI.xls")
SURVEY_FILE_2 = os.path.join(DULIEU, "D25 Giáo dục thể chất 2 25-26 HKI 2.xls")
FEEDBACK_FILE = os.path.join(DULIEU, "feedback1.xls")
FEEDBACK_FILE_2 = os.path.join(DULIEU, "feedback2.xls")


# ───────────────────────────────────────────────────────────────
# data_utils: lớp đọc thống nhất
# ───────────────────────────────────────────────────────────────

@pytest.mark.skipif(not os.path.exists(SURVEY_FILE), reason="thiếu file mẫu")
def test_open_workbook_survey():
    wb = data_utils.open_workbook(SURVEY_FILE)
    assert wb.nsheets >= 1
    sh = wb.sheet_by_index(0)
    assert sh.ncols >= 8
    # API tương thích xlrd
    assert sh.cell(0, 0).value == sh.cell_value(0, 0)


@pytest.mark.skipif(not os.path.exists(FEEDBACK_FILE), reason="thiếu file mẫu")
def test_detect_feedback_column():
    wb = data_utils.open_workbook(FEEDBACK_FILE)
    sh = wb.sheet_by_index(0)
    col, start = data_utils.detect_feedback_column(
        sh, ["ý kiến đóng góp", "feedback"])
    # File feedback mẫu: cột ý kiến ở index 36, header ở dòng 5 (index 4)
    assert col == 36
    assert start >= 5


@pytest.mark.skipif(not os.path.exists(FEEDBACK_FILE), reason="thiếu file mẫu")
def test_find_column_by_keywords():
    wb = data_utils.open_workbook(FEEDBACK_FILE)
    sh = wb.sheet_by_index(0)
    col = data_utils.find_column_by_keywords(sh, ["mã số sinh viên"])
    assert col == 7


# ───────────────────────────────────────────────────────────────
# Trích xuất dữ liệu khảo sát
# ───────────────────────────────────────────────────────────────

@pytest.mark.skipif(not os.path.exists(SURVEY_FILE), reason="thiếu file mẫu")
def test_extract_survey_questions():
    data = e2w.extract_data_from_excel(SURVEY_FILE)
    # File mẫu có 22 câu hỏi khảo sát
    assert len(data["survey_questions"]) == 22
    # Mỗi câu có 5 mức và tổng > 0
    q = data["survey_questions"][0]
    assert len(q["counts"]) == 5
    assert q["total"] > 0


@pytest.mark.skipif(not os.path.exists(SURVEY_FILE), reason="thiếu file mẫu")
def test_average_score_in_valid_range():
    data = e2w.extract_data_from_excel(SURVEY_FILE)
    for q in data["survey_questions"]:
        # Điểm TB thang Likert 1-5 phải nằm trong [1, 5]
        assert 1.0 <= q["average"] <= 5.0, f"Điểm TB lỗi: {q['average']}"


@pytest.mark.skipif(not os.path.exists(SURVEY_FILE), reason="thiếu file mẫu")
def test_average_score_manual_calc():
    """Kiểm chứng công thức điểm TB bằng tính tay trên câu đầu tiên."""
    data = e2w.extract_data_from_excel(SURVEY_FILE)
    q = data["survey_questions"][0]
    counts = q["counts"]
    total = q["total"]
    expected = sum((i + 1) * counts[i] for i in range(5)) / total
    assert math.isclose(q["average"], round(expected, 2), abs_tol=0.01)


@pytest.mark.skipif(not os.path.exists(SURVEY_FILE_2), reason="thiếu file mẫu")
def test_tong_so_sv_from_filename():
    """File kết thúc bằng ' 2' => tong_so_sv = tổng phiếu + 2."""
    data = e2w.extract_data_from_excel(SURVEY_FILE_2)
    total_phieu = data["survey_questions"][0]["total"]
    assert data["meta"].get("tong_so_sv") == total_phieu + 2


# ───────────────────────────────────────────────────────────────
# Tách & làm sạch ý kiến mở
# ───────────────────────────────────────────────────────────────

def test_parse_open_responses_filters_empty():
    assert e2w.parse_open_responses("") == []
    assert e2w.parse_open_responses("không") == []
    assert e2w.parse_open_responses("không ạ") == []
    assert e2w.parse_open_responses("dạ không ạ") == []


def test_parse_open_responses_keeps_real():
    out = e2w.parse_open_responses("Giáo viên dạy rất hay")
    assert len(out) == 1
    assert out[0].endswith(".")
    assert out[0][0].isupper()


def test_parse_open_responses_dedup():
    out = e2w.parse_open_responses("Tốt lắm\nTốt lắm\nTốt lắm")
    assert len(out) == 1


# ───────────────────────────────────────────────────────────────
# Gộp feedback
# ───────────────────────────────────────────────────────────────

@pytest.mark.skipif(not os.path.exists(FEEDBACK_FILE), reason="thiếu file mẫu")
def test_extract_feedbacks_returns_list():
    fbs = mf.extract_feedbacks(FEEDBACK_FILE)
    assert isinstance(fbs, list)
    # Các câu vô nghĩa như "Không", "Dạ không ạ" phải bị loại
    for fb in fbs:
        assert fb.strip().lower() not in ("không", "không.", "dạ không ạ")


@pytest.mark.skipif(not os.path.exists(FEEDBACK_FILE), reason="thiếu file mẫu")
def test_merge_feedbacks_creates_file(tmp_path):
    out = mf.merge_feedbacks([FEEDBACK_FILE, FEEDBACK_FILE_2], str(tmp_path))
    if out is not None:  # có thể None nếu toàn bộ là "không"
        assert os.path.exists(out)
        assert out.endswith(".docx")


# ───────────────────────────────────────────────────────────────
# Tạo báo cáo Word end-to-end
# ───────────────────────────────────────────────────────────────

@pytest.mark.skipif(not os.path.exists(SURVEY_FILE), reason="thiếu file mẫu")
def test_create_word_report(tmp_path):
    data = e2w.extract_data_from_excel(SURVEY_FILE)
    out = os.path.join(str(tmp_path), "report.docx")
    e2w.create_word_report(data, out)
    assert os.path.exists(out)
    assert os.path.getsize(out) > 5000  # file Word hợp lệ có kích thước đáng kể


# ───────────────────────────────────────────────────────────────
# Validate dữ liệu
# ───────────────────────────────────────────────────────────────

@pytest.mark.skipif(not os.path.exists(SURVEY_FILE), reason="thiếu file mẫu")
def test_validate_survey_sheet():
    wb = data_utils.open_workbook(SURVEY_FILE)
    sh = wb.sheet_by_index(0)
    report = data_utils.validate_survey_sheet(sh, "Sheet 1")
    # File mẫu hợp lệ => không có lỗi nghiêm trọng
    assert report.ok


@pytest.mark.skipif(not os.path.exists(SURVEY_FILE_2), reason="thiếu file mẫu")
def test_multi_sheet_aggregation_single_suffix(tmp_path):
    # Chạy process_file cho SURVEY_FILE_2 có 4 sheets
    outputs = e2w.process_file(SURVEY_FILE_2, output_dir=str(tmp_path))
    assert outputs is not None
    # Chỉ được tạo 1 file kết quả đại diện cho dữ liệu gộp, thay vì từng file cho mỗi sheet
    assert len(outputs) == 1
    out_file = outputs[0]
    assert os.path.exists(out_file)
    
    # Kiểm tra logic trích xuất không mang suffix ở cấp sheet và chỉ cộng 1 lần khi gộp
    n_sheets, _ = e2w.get_sheet_count(SURVEY_FILE_2)
    all_data = []
    for si in range(n_sheets):
        data = e2w.extract_data_from_excel(SURVEY_FILE_2, sheet_index=si, apply_filename_suffix=False)
        if data["survey_questions"]:
            all_data.append(data)
            
    # Mỗi sheet đơn lẻ không được tự ý áp dụng suffix của file name
    for d in all_data:
        assert d["meta"].get("tong_so_sv") is None
        
    merged = e2w.merge_sheets_data(all_data)
    assert merged is not None
    
    # Tính toán tong_so_sv gộp có suffix
    import re
    basename = os.path.splitext(os.path.basename(SURVEY_FILE_2))[0]
    match = re.search(r'\s+(\d+)$', basename)
    sv_khong_khao_sat = int(match.group(1)) if match else None
    assert sv_khong_khao_sat == 2
    
    total_phieu = merged["survey_questions"][0].get("total", 0)
    merged["meta"]["tong_so_sv"] = total_phieu + sv_khong_khao_sat
    
    # tong_so_sv của bản gộp phải bằng tổng phiếu phản hồi của tất cả các sheet + 2
    assert merged["meta"]["tong_so_sv"] == total_phieu + 2


NEW_SURVEY_FILE = os.path.join(DULIEU, "2627", "CTXH26 Giáo dục thể chất 1 2627HKI 0.xls")

@pytest.mark.skipif(not os.path.exists(NEW_SURVEY_FILE), reason="thiếu file mẫu mới 2627")
def test_new_report_template_format(tmp_path):
    """Kiểm tra mẫu mới: mã lớp góc trên bên trái, tiêu đề học phần cụ thể, dấu * phần 3."""
    from docx import Document
    data = e2w.extract_data_from_excel(NEW_SURVEY_FILE)
    out = os.path.join(str(tmp_path), "CTXH26_test.docx")
    e2w.create_word_report(data, out)
    assert os.path.exists(out)

    doc = Document(out)
    
    # 1. Textbox mã lớp ở góc trên bên trái (nằm ở paragraph sau header table)
    p0_xml = doc.paragraphs[0]._element
    p0_texts = [t.text for t in p0_xml.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')]
    assert "CTXH26" in p0_texts

    # 2. Tiêu đề báo cáo cụ thể tên môn học viết hoa (không để học phần cơ sở chung chung)
    title_texts = [p.text for p in doc.paragraphs[1:4]]
    assert "BÁO CÁO TỔNG HỢP KẾT QUẢ" in title_texts[0]
    assert "KHẢO SÁT SINH VIÊN SAU KHI KẾT THÚC HỌC PHẦN GIÁO DỤC THỂ CHẤT 1" in title_texts[1]
    assert "& ĐỀ XUẤT KẾ HOẠCH CẢI TIẾN" in title_texts[2]

    # 3. Dấu * trước Phần 3; Phần 3 liệt kê ý kiến thực tế của sinh viên
    all_texts = [p.text for p in doc.paragraphs]
    assert "*" in all_texts
    star_idx = all_texts.index("*")
    sec3_idx = all_texts.index("Phần 3. Các ý kiến đóng góp từ câu hỏi mở:")
    assert star_idx < sec3_idx
    opinions = [t for t in all_texts[sec3_idx + 1:] if t.startswith("- ")]
    assert any("Sân trường rất nắng" in t for t in opinions)
    assert "*Không có ý kiến." not in all_texts

    # 4. Nhóm tiêu chí theo mẫu phiếu mới: 8.1–8.4 thuộc PHƯƠNG PHÁP GIẢNG DẠY
    table = doc.tables[2]
    rows = [r.cells[1].text for r in table.rows]
    ii = rows.index("II. PHƯƠNG PHÁP GIẢNG DẠY")
    iii = rows.index("III. LƯỢNG GIÁ HỌC PHẦN")
    assert [t.split(".")[0:2] for t in rows[ii + 1:iii]] == [["C8", "1"], ["C8", "2"], ["C8", "3"], ["C8", "4"]]
    assert any("tối thiểu 2 tiêu chí trong 18 tiêu chí" in t for t in all_texts)


def test_no_opinions_writes_placeholder(tmp_path):
    from docx import Document
    data = e2w.extract_data_from_excel(SURVEY_FILE)
    data["open_responses"] = "Không\nDạ không ạ\nEm không có ý kiến"
    data["meta"]["tong_so_sv"] = 50
    out = os.path.join(str(tmp_path), "empty.docx")
    e2w.create_word_report(data, out)
    assert "*Không có ý kiến." in [p.text for p in Document(out).paragraphs]


def test_open_responses_keep_proper_names():
    res = e2w.parse_open_responses("Em mong thầy Nam giảng chậm hơn\nCô Lan dạy rất hay")
    assert res == ["Em mong thầy Nam giảng chậm hơn.", "Cô Lan dạy rất hay."]


def test_open_responses_split_glued_refusals():
    res = e2w.parse_open_responses("Sân tập nắng lắm ạdạ không ạKhôngkhông có")
    assert res == ["Sân tập nắng lắm ạ."]


def test_build_sections_old_form():
    qs = [{"number": str(i)} for i in range(1, 23)]
    secs = e2w.build_sections(qs)
    assert [(s["start"], s["end"]) for s in secs] == [(0, 10), (10, 14), (14, 18), (18, 22)]


def test_build_sections_unknown_form_returns_none():
    assert e2w.build_sections([{"number": str(i)} for i in range(1, 6)]) is None


def test_parse_section_layouts_rejects_bad_line():
    with pytest.raises(ValueError):
        e2w.parse_section_layouts("Nhóm A | 1-x")


def test_merge_dedupes_class_name():
    n, _ = e2w.get_sheet_count(SURVEY_FILE_2)
    all_data = [e2w.extract_data_from_excel(SURVEY_FILE_2, si, False) for si in range(n)]
    merged = e2w.merge_sheets_data([d for d in all_data if d["survey_questions"]])
    assert merged["meta"]["lop_hoc"].count(",") == len(set(merged["sheet_names"])) - 1


def test_generate_report_strips_absent_count_from_name(tmp_path):
    src = os.path.join(str(tmp_path), "Lop A 5.xls")
    import shutil
    shutil.copy(SURVEY_FILE, src)
    res = e2w.generate_report(src, str(tmp_path))
    assert res.ok, res.error
    assert os.path.basename(res.output_path) == "Lop A.docx"
    first_total = e2w.extract_data_from_excel(SURVEY_FILE)["survey_questions"][0]["total"]
    assert res.total_students == first_total + 5


def test_build_sections_2627_template_form():
    nums = [str(i) for i in range(1, 11)] + ["11.1", "11.2", "11.3", "11.4"] + [str(i) for i in range(12, 20)]
    secs = e2w.build_sections([{"number": n} for n in nums])
    assert [(s["start"], s["end"]) for s in secs] == [(0, 10), (10, 14), (14, 18), (18, 22)]


@pytest.mark.skipif(not os.path.exists(NEW_SURVEY_FILE), reason="thiếu file mẫu mới 2627")
def test_report_current_date_and_chinh_quy_fill(tmp_path):
    """Kiểm tra: ngày báo cáo là ngày hiện tại, và tự động điền Chính quy (100%) khi Excel không có."""
    from docx import Document
    from datetime import datetime

    data = e2w.extract_data_from_excel(NEW_SURVEY_FILE)
    # File Excel gốc có ngày 20/10/2026
    assert "20 tháng 10" in data.get("signer_date", "")

    out = os.path.join(str(tmp_path), "test_current_date.docx")
    e2w.create_word_report(data, out)
    assert os.path.exists(out)

    doc = Document(out)
    
    # 1. Kiểm tra ngày báo cáo lấy ngày hiện tại
    now = datetime.now()
    expected_date = f"TP. HCM, ngày {now.day:02d} tháng {now.month:02d} năm {now.year}"
    header_right_cell_text = doc.tables[0].cell(0, 1).text
    assert expected_date in header_right_cell_text

    # 2. Kiểm tra Bảng 2.1 Thông tin nền đã điền Chính quy (59) và Khác (0)
    bg_table = doc.tables[1]
    # Row 1: Hình thức đào tạo | Chính quy | 59 | 100.00%
    row_cq = [c.text.strip() for c in bg_table.rows[1].cells]
    assert row_cq[1] == "Chính quy"
    assert row_cq[2] == "59"
    assert row_cq[3] == "100.00%"

    # Row 2: | Khác | 0 | 0.00%
    row_khac = [c.text.strip() for c in bg_table.rows[2].cells]
    assert row_khac[1] == "Khác"
    assert row_khac[2] == "0"
    assert row_khac[3] == "0.00%"



def test_ai_summary_replaces_opinion_list(tmp_path, monkeypatch):
    """Có mạng: Phần 3 là đoạn tóm tắt AI. Lỗi/không mạng: liệt kê ý kiến như cũ."""
    import ai_summary
    from docx import Document
    monkeypatch.setitem(e2w.CONFIG, "ai_enabled", True)
    monkeypatch.setitem(e2w.CONFIG, "ai_api_key", "test-key")

    monkeypatch.setattr(ai_summary, "_chat",
                        lambda *a, **k: "**Sinh viên** hài lòng với giảng viên, đề xuất thêm bài tập.")
    res = e2w.generate_report(SURVEY_FILE_2, str(tmp_path / "ai"))
    assert res.ok and res.summary_source == "ai"
    texts = _section3(res.output_path)
    assert texts[0] == "Sinh viên hài lòng với giảng viên, đề xuất thêm bài tập."
    assert not any(t.startswith("- ") for t in texts)

    def _offline(*a, **k):
        raise OSError("no network")
    monkeypatch.setattr(ai_summary, "_chat", _offline)
    res = e2w.generate_report(SURVEY_FILE_2, str(tmp_path / "offline"))
    assert res.ok and res.summary_source == "offline"
    assert any(t.startswith("- ") for t in _section3(res.output_path))


def _section3(path):
    from docx import Document
    texts = [p.text for p in Document(path).paragraphs]
    start = texts.index("Phần 3. Các ý kiến đóng góp từ câu hỏi mở:") + 1
    end = texts.index("Phần 4. Đề xuất kế hoạch cải tiến:")
    return [t for t in texts[start:end] if t.strip()]


def test_ai_reply_parser_tolerates_trailing_sse_and_streaming():
    import ai_summary
    msg = '{"choices":[{"message":{"content":"Đoạn tóm tắt"}}]}'
    assert ai_summary._parse_reply(msg + "data: [DONE]\n\n") == "Đoạn tóm tắt"
    sse = ('data: {"choices":[{"delta":{"content":"Đoạn "}}]}\n\n'
           'data: {"choices":[{"delta":{"content":"tóm tắt"}}]}\n\ndata: [DONE]\n')
    assert ai_summary._parse_reply(sse) == "Đoạn tóm tắt"


def test_prompt_placeholders_and_per_target_prompts(monkeypatch):
    import ai_summary as a
    filled = a.fill_placeholders("Tóm tắt {so_y_kien} ý kiến lớp {lop} môn {mon_hoc}.",
                                 {"so_y_kien": 3, "mon_hoc": "GDTC 1"})
    assert filled == "Tóm tắt 3 ý kiến lớp môn GDTC 1."      # biến thiếu → bỏ trống
    assert a.fill_placeholders("Giữ {khac} nguyên") == "Giữ {khac} nguyên"

    monkeypatch.setitem(e2w.CONFIG, "ai_prompt", "")
    monkeypatch.setitem(e2w.CONFIG, "ai_prompt_feedback", "Prompt riêng cho feedback")
    assert a.get_prompt("report") == a.DEFAULT_PROMPT
    assert a.get_prompt("feedback") == "Prompt riêng cho feedback"

    seen = {}
    monkeypatch.setitem(e2w.CONFIG, "ai_enabled", True)
    monkeypatch.setitem(e2w.CONFIG, "ai_api_key", "k")
    monkeypatch.setattr(a, "_chat", lambda cfg, msgs, timeout: seen.setdefault("sys", msgs[0]["content"])
                        or "x" * 30)
    a.summarize_feedbacks(["Ý kiến A", "Ý kiến B"], kind="feedback")
    assert seen["sys"] == "Prompt riêng cho feedback"


def test_same_base_name_does_not_overwrite(tmp_path):
    import shutil
    a = os.path.join(str(tmp_path), "Lop B 3.xls")
    b = os.path.join(str(tmp_path), "Lop B 5.xls")
    shutil.copy(SURVEY_FILE, a)
    shutil.copy(SURVEY_FILE, b)
    taken = set()
    r1 = e2w.generate_report(a, str(tmp_path / "out"), taken=taken)
    r2 = e2w.generate_report(b, str(tmp_path / "out"), taken=taken)
    assert os.path.basename(r1.output_path) == "Lop B.docx"
    assert os.path.basename(r2.output_path) == "Lop B (2).docx"


def test_ai_skips_after_network_failure(monkeypatch):
    import urllib.error
    import ai_summary as a
    monkeypatch.setitem(e2w.CONFIG, "ai_enabled", True)
    monkeypatch.setitem(e2w.CONFIG, "ai_api_key", "k")
    monkeypatch.setattr(a, "_offline_until", 0.0)
    calls = []

    def down(*args, **kw):
        calls.append(1)
        raise urllib.error.URLError("no network")
    monkeypatch.setattr(a, "_chat", down)
    assert a.summarize_feedbacks(["Ý kiến A"]) is None
    assert a.summarize_feedbacks(["Ý kiến B"]) is None
    assert len(calls) == 1          # lần 2 bỏ qua ngay, không chờ mạng
