#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tạo BÁO CÁO PHÂN TÍCH (companion report) kèm biểu đồ cho SKKN (Round 2).

Khác với báo cáo Word chính thức theo mẫu ĐBCL (excel_to_word.create_word_report),
module này tạo MỘT file phân tích riêng "Bao_Cao_Phan_Tich.docx" gồm:
  - Bảng chỉ số tổng quan (điểm TB, SD, % hài lòng).
  - Biểu đồ phân bố, điểm TB theo tiêu chí, donut hài lòng.
  - Bảng gợi ý tiêu chí cần cải tiến (điểm thấp nhất).
  - (Tùy chọn) biểu đồ xu hướng giữa các kỳ.

Giữ tách biệt để KHÔNG ảnh hưởng định dạng báo cáo nộp chính thức.
"""

import os
from typing import List, Optional, Dict

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

from analytics import SurveyAnalysis, suggest_improvements
import charts

FONT = "Times New Roman"
HEADER_BG = "D9E2F3"


def _set_cell_bg(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def _set_run(run, size=12, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.name = FONT
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _para(doc, text, size=12, bold=False, italic=False, align=None,
          space_after=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run(r, size, bold, italic, color)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def _cell_text(cell, text, size=11, bold=False, align=None, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    _set_run(r, size, bold, color=color)
    if align is not None:
        p.alignment = align


def _setup_page(doc):
    s = doc.sections[0]
    s.page_width = Cm(21)
    s.page_height = Cm(29.7)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2)
    s.top_margin = Cm(2)
    s.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5


def _add_overview_table(doc, analysis: SurveyAnalysis):
    rows = [
        ("Học phần / Lớp", analysis.label),
        ("Số tiêu chí đánh giá", str(analysis.n_questions)),
        ("Số phiếu phản hồi", str(analysis.n_responses)),
        ("Điểm trung bình chung", f"{analysis.overall_average:.2f} / 5"),
        ("Độ lệch chuẩn (SD)", f"{analysis.overall_sd:.2f}"),
        ("Tỷ lệ hài lòng (mức 4-5)", f"{analysis.overall_satisfied_pct:.1f}%"),
    ]
    if analysis.highest:
        rows.append(("Tiêu chí điểm cao nhất",
                     f"C{analysis.highest.number} ({analysis.highest.average:.2f})"))
    if analysis.lowest:
        rows.append(("Tiêu chí điểm thấp nhất",
                     f"C{analysis.lowest.number} ({analysis.lowest.average:.2f})"))

    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        _cell_text(table.cell(i, 0), k, size=11, bold=True)
        _set_cell_bg(table.cell(i, 0), HEADER_BG)
        _cell_text(table.cell(i, 1), v, size=11)
        table.cell(i, 0).width = Cm(7)
        table.cell(i, 1).width = Cm(9)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0


def _add_image(doc, path, width_in=6.0):
    if path and os.path.exists(path):
        doc.add_picture(path, width=Inches(width_in))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_improvement_table(doc, analysis: SurveyAnalysis, top_n=3):
    suggestions = suggest_improvements(analysis, top_n=top_n)
    if not suggestions:
        return
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Tiêu chí", "Điểm TB", "% chưa hài lòng", "Nội dung tiêu chí"]
    for i, h in enumerate(headers):
        _cell_text(table.cell(0, i), h, size=11, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_bg(table.cell(0, i), HEADER_BG)
    for q in suggestions:
        row = table.add_row()
        _cell_text(row.cells[0], f"C{q.number}", size=11,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(row.cells[1], f"{q.average:.2f}", size=11,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(row.cells[2], f"{q.dissatisfied_pct:.1f}%", size=11,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(row.cells[3], q.text, size=10)
    for i, w in enumerate([2, 2, 3, 9]):
        for row in table.rows:
            row.cells[i].width = Cm(w)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0


def create_analysis_report(analysis: SurveyAnalysis, output_path: str,
                           chart_dir: Optional[str] = None,
                           history: Optional[List[Dict]] = None) -> str:
    """Tạo file Word phân tích kèm biểu đồ. Trả về đường dẫn file."""
    out_dir = os.path.dirname(output_path) or "."
    os.makedirs(out_dir, exist_ok=True)
    chart_dir = chart_dir or os.path.join(out_dir, "_charts")

    # Sinh biểu đồ
    chart_files = charts.generate_all_charts(analysis, chart_dir, prefix=analysis.label)
    trend_path = None
    if history and len(history) >= 2:
        trend_path = charts.chart_trend(
            history, os.path.join(chart_dir, "trend.png"))

    doc = Document()
    _setup_page(doc)

    _para(doc, "BÁO CÁO PHÂN TÍCH KẾT QUẢ KHẢO SÁT", size=15, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _para(doc, f"({analysis.label})", size=12, italic=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    _para(doc, "1. Tổng quan chỉ số", size=13, bold=True, space_after=6)
    _add_overview_table(doc, analysis)
    doc.add_paragraph()

    _para(doc, "2. Tỷ lệ hài lòng", size=13, bold=True, space_after=6)
    _add_image(doc, chart_files.get("satisfaction"), width_in=4.5)
    doc.add_paragraph()

    _para(doc, "3. Phân bố mức độ đồng ý", size=13, bold=True, space_after=6)
    _add_image(doc, chart_files.get("distribution"), width_in=6.0)
    doc.add_paragraph()

    _para(doc, "4. Điểm trung bình theo từng tiêu chí", size=13, bold=True,
          space_after=6)
    _add_image(doc, chart_files.get("averages"), width_in=6.2)
    doc.add_paragraph()

    if trend_path:
        _para(doc, "5. Xu hướng qua các kỳ khảo sát", size=13, bold=True,
              space_after=6)
        _add_image(doc, trend_path, width_in=6.2)
        doc.add_paragraph()

    sec = 6 if trend_path else 5
    _para(doc, f"{sec}. Gợi ý tiêu chí cần ưu tiên cải tiến", size=13,
          bold=True, space_after=2)
    _para(doc, "(Các tiêu chí có điểm trung bình thấp nhất - đề xuất tự động)",
          size=10, italic=True, space_after=6)
    _add_improvement_table(doc, analysis)

    doc.save(output_path)
    return output_path
