#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Module gộp và phân loại ý kiến đóng góp mở từ nhiều file khảo sát Excel (.xls / .xlsx).
"""
import sys
import os
import argparse
from typing import Callable, List, Optional
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from excel_to_word import parse_open_responses, CONFIG
from data_utils import open_workbook, detect_feedback_column
from app_logging import get_logger
import feedback_classifier as fc

log = get_logger()


def _is_survey_summary(sheet) -> bool:
    """Sheet thống kê kết quả khảo sát (có mục "Thông tin khảo sát" ở cột A)."""
    return any(str(sheet.cell_value(r, 0)).strip() == "Thông tin khảo sát"
               for r in range(min(80, sheet.nrows)))


def extract_feedbacks(filepath: str, raise_errors: bool = False) -> List[str]:
    """Trích xuất ý kiến mở từ file feedback (.xls/.xlsx).

    Dò cột feedback TỰ ĐỘNG theo tiêu đề (không hard-code index 36),
    có fallback an toàn về cột cuối. Ghi log số ý kiến tìm được.
    """
    feedbacks = []
    try:
        wb = open_workbook(filepath)
        search_keywords = CONFIG.get("feedback_search_keywords", ["ý kiến đóng góp", "feedback"])
        skip_keywords = CONFIG.get("feedback_skip_keywords",
                                   ["không", "không có", "khong", "none", "n/a",
                                    "chưa có", "không có ý kiến", "không có ý kiến gì"])
        skip_keywords_lower = [k.lower() for k in skip_keywords]

        for si, sheet in enumerate(wb.sheets):
            if sheet.nrows == 0:
                continue

            # File THỐNG KÊ khảo sát (không phải file ý kiến thô): lấy đúng ô câu hỏi mở,
            # tránh đọc nhầm cột số liệu ("133.0", "Tổng cộng") thành ý kiến.
            if _is_survey_summary(sheet):
                from excel_to_word import extract_data_from_excel
                data = extract_data_from_excel(filepath, sheet_index=si, apply_filename_suffix=False)
                responses = parse_open_responses(data.get("open_responses", ""))
                feedbacks.extend(responses)
                log.info("Đọc ý kiến mở từ file thống kê '%s' [%s]: %d ý kiến",
                         os.path.basename(filepath), sheet.name, len(responses))
                continue

            feedback_col_idx, start_row = detect_feedback_column(sheet, search_keywords)
            log.info("Đọc feedback '%s' [%s]: cột %d, từ dòng %d",
                     os.path.basename(filepath), sheet.name, feedback_col_idx, start_row + 1)

            sheet_count = 0
            for r in range(start_row, sheet.nrows):
                val = str(sheet.cell_value(r, feedback_col_idx)).strip()
                # Loại bỏ các câu trả lời rỗng hoặc không có ý nghĩa
                if val and val.lower() not in skip_keywords_lower:
                    cleaned_responses = parse_open_responses(val)
                    feedbacks.extend(cleaned_responses)
                    sheet_count += len(cleaned_responses)
            log.info("  -> %d ý kiến hợp lệ từ sheet '%s'", sheet_count, sheet.name)
    except Exception as e:
        log.error("Lỗi đọc file %s: %s", filepath, e)
        if raise_errors:
            raise
        print(f"  ❌ Lỗi đọc file {os.path.basename(filepath)}: {e}")
    return feedbacks


def merge_feedbacks(input_files: List[str], output_dir: Optional[str] = None,
                    on_progress: Optional[Callable[[str, str], None]] = None) -> Optional[str]:
    """Gộp ý kiến mở từ nhiều file khảo sát, phân loại theo chủ đề/sắc thái và xuất ra Word.

    on_progress(msg, level): nhận thông báo từng file ("info" | "warning" | "error").
    Trả về đường dẫn Feedbacks.docx, hoặc None nếu không có ý kiến hợp lệ.
    Ném PermissionError (thông điệp tiếng Việt) nếu không ghi được file kết quả.
    """
    def notify(msg, level="info"):
        print(f"  {msg}")
        if on_progress:
            on_progress(msg, level)

    if not input_files:
        return None
    
    out_dir = output_dir or os.path.dirname(input_files[0])
    os.makedirs(out_dir, exist_ok=True)
    output_path = os.path.join(out_dir, "Feedbacks.docx")
    
    all_feedbacks = []
    for filepath in input_files:
        name = os.path.basename(filepath)
        try:
            fb = extract_feedbacks(filepath, raise_errors=True)
        except Exception as e:
            notify(f"❌ {name}: {e}", "error")
            continue
        notify(f"📄 {name}: {len(fb)} ý kiến", "info" if fb else "warning")
        all_feedbacks.extend(fb)
        
    # Loại bỏ trùng lặp nếu cần
    seen = set()
    unique_feedbacks = []
    for fb in all_feedbacks:
        fb_lower = fb.lower().strip()
        if fb_lower not in seen:
            seen.add(fb_lower)
            unique_feedbacks.append(fb)
            
    if not unique_feedbacks:
        print("  ⚠️ Không tìm thấy feedback nào hợp lệ.")
        return None
        
    # Tạo file Word
    doc = Document()
    
    # Cấu hình trang
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    
    # Style
    style = doc.styles['Normal']
    font = style.font
    font.name = CONFIG["font_name"]
    font.size = Pt(CONFIG["font_size"])
    style.paragraph_format.line_spacing = 1.5
    
    heading = doc.add_paragraph("TỔNG HỢP FEEDBACKS", style='Heading 1')
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Tổng số ý kiến đóng góp: {len(unique_feedbacks)}")

    # ── TÓM TẮT BẰNG AI (khi có mạng; lỗi/offline thì bỏ qua phần này) ──
    import ai_summary
    summary = ai_summary.summarize_feedbacks(unique_feedbacks, kind="feedback")
    if summary:
        doc.add_paragraph()
        doc.add_paragraph("TÓM TẮT ĐÁNH GIÁ CỦA SINH VIÊN", style='Heading 2')
        p_sum = doc.add_paragraph(summary)
        p_sum.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        notify("🤖 Đã tóm tắt ý kiến bằng AI")
    elif ai_summary.is_configured():
        notify("📴 Không kết nối được AI — chỉ phân loại offline như cũ", "warning")

    # ── PHÂN LOẠI Ý KIẾN THEO CHỦ ĐỀ & SẮC THÁI (offline) ──
    try:
        result = fc.classify_feedbacks(unique_feedbacks)
        log.info("Phân loại %d ý kiến theo chủ đề & sắc thái", result.total)

        doc.add_paragraph()
        doc.add_paragraph("PHÂN LOẠI THEO CHỦ ĐỀ", style='Heading 2')
        for topic, pct in result.topic_percentages().items():
            cnt = result.topic_counts[topic]
            doc.add_paragraph(f"• {topic}: {cnt} ý kiến ({pct}%)")

        doc.add_paragraph()
        doc.add_paragraph("PHÂN LOẠI THEO SẮC THÁI", style='Heading 2')
        for s, pct in result.sentiment_percentages().items():
            cnt = result.sentiment_counts[s]
            doc.add_paragraph(f"• {s.capitalize()}: {cnt} ý kiến ({pct}%)")

        # Liệt kê ý kiến gom theo chủ đề
        doc.add_paragraph()
        doc.add_paragraph("CHI TIẾT Ý KIẾN THEO CHỦ ĐỀ", style='Heading 2')
        for topic, items in result.by_topic().items():
            doc.add_paragraph()
            ph = doc.add_paragraph(f"{topic} ({len(items)} ý kiến)")
            ph.runs[0].bold = True
            for i, it in enumerate(items, 1):
                p = doc.add_paragraph(f"{i}. [{it.sentiment}] {it.text}")
                p.paragraph_format.space_after = Pt(4)
    except Exception as e:
        log.warning("Không phân loại được ý kiến: %s", e)
        # Fallback: liệt kê đơn giản như cũ
        doc.add_paragraph()
        for i, fb in enumerate(unique_feedbacks, 1):
            p = doc.add_paragraph(f"{i}. {fb}")
            p.paragraph_format.space_after = Pt(6)

    try:
        doc.save(output_path)
    except PermissionError:
        raise PermissionError(f"Không ghi được “{os.path.basename(output_path)}” — file đang mở "
                              "trong Word? Hãy đóng file rồi chạy lại") from None
    log.info("Đã lưu %d feedbacks vào %s", len(unique_feedbacks), os.path.basename(output_path))
    print(f"  ✅ Đã lưu {len(unique_feedbacks)} feedbacks vào {os.path.basename(output_path)}")
    return output_path


def main():
    """Hàm chạy dòng lệnh (CLI)."""
    from excel_to_word import _safe_console
    _safe_console()
    parser = argparse.ArgumentParser(description="Gộp và phân loại feedbacks từ nhiều file Excel khảo sát")
    parser.add_argument("files", nargs="*", help="Danh sách file Excel (.xls/.xlsx)")
    parser.add_argument("--output-dir", "-o", help="Thư mục lưu file Feedbacks.docx kết quả")
    args = parser.parse_args()
    
    files = args.files if args.files else [f for f in os.listdir('.') if f.endswith(('.xls', '.xlsx'))]
    merge_feedbacks(files, output_dir=args.output_dir)


if __name__ == "__main__":
    main()
